#!/usr/bin/env python3
"""Render the four legacy PRP1 oficio fixtures from in-house transcriptions.

The cases 222AAA / 333BBB / 333ccc / 555CCC are long-lived test fixtures. This script
re-creates them from the synthetic markdown sources in this folder so every PDF, page
image and DOCX in the repo is produced by our own tooling (neutral header, a
"DOCUMENTO SINTÉTICO" watermark, and no third-party metadata).

Outputs (per case) into --out:
  <case>.pdf                 image-only "scan" (scanned cases) or text PDF (digital cases)
  <case>_page*.{jpg,png}     page rasters with the historical names/sizes the tests expect
  <case>.docx                CNBV cover letter rebuilt from --docx-src text (optional)
  <case>_page-0001.ocr.txt   Tesseract text for 222AAA page 1 (when --tesseract is given)

Requires: chromium/google-chrome, pdftoppm (poppler), Pillow, numpy, markdown, python-docx.
Usage:
  python render_legacy_oficios.py --out OUT [--docx-src DIR] [--tesseract BIN --tessdata DIR]
"""
from __future__ import annotations

import argparse
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import markdown
import numpy as np
from PIL import Image, ImageFilter

HERE = Path(__file__).resolve().parent
# Optional text substitutions applied when rebuilding cover letters (old text -> new text).
DOCX_TEXT_SUBS: dict[str, str] = {}
WATERMARK = "DOCUMENTO SINTÉTICO"

# case -> (scanned?, raster outputs as (name, width, height, format))
CASES: dict[str, tuple[bool, list[tuple[str, int, int, str]]]] = {
    "222AAA-44444444442025": (True, [
        ("222AAA-44444444442025_page-0001.jpg", 1275, 1650, "JPEG"),
        ("222AAA-44444444442025_page-0002.jpg", 1275, 1650, "JPEG"),
        ("222AAA-44444444442025_page-0003.jpg", 1275, 1650, "JPEG"),
        ("222AAA-44444444442025_page-0004.jpg", 1275, 1650, "JPEG"),
        ("222AAA-44444444442025_page-1.png", 2550, 3300, "PNG"),
    ]),
    "333BBB-44444444442025": (True, [("333BBB-44444444442025_page1.png", 1020, 1320, "PNG")]),
    "333ccc-6666666662025": (True, [("333ccc-6666666662025_page1.png", 1020, 1320, "PNG")]),
    "555CCC-66666662025": (False, [
        ("555CCC-66666662025_page1.png", 1020, 1320, "PNG"),
        ("555CCC-66666662025_page-0001.png", 1020, 1320, "PNG"),
    ]),
}

HEADER_SVG = """<svg xmlns="http://www.w3.org/2000/svg" width="520" height="54" viewBox="0 0 520 54">
<rect x="0" y="0" width="520" height="54" fill="#e9e9e9"/>
<circle cx="30" cy="27" r="18" fill="none" stroke="#555" stroke-width="3"/>
<path d="M20 33 L30 15 L40 33 Z" fill="#555"/>
<text x="60" y="23" font-family="Arial" font-size="13" fill="#333">EXXERCUBE PRISMA · CORPUS SINTÉTICO</text>
<text x="60" y="42" font-family="Arial" font-size="11" fill="#555">Formato de prueba de software — no es un documento oficial</text>
</svg>"""

CSS = """
@page { size: Letter; margin: 36px 48px; }
body { font-family: Arial, Helvetica, sans-serif; font-size: 11pt; line-height: 1.35; color: #000; }
.page { page-break-after: always; position: relative; }
.page:last-child { page-break-after: auto; }
.hdr { text-align: center; margin-bottom: 8px; }
.id-box { border: 1px solid #000; padding: 4px 10px; float: right; margin: 8px 0 14px 20px; font-size: 10pt; }
.addressee { clear: both; font-size: 13pt; line-height: 1.5; margin-top: 14px; }
.spaced { letter-spacing: 4px; font-size: 10pt; }
.section-header { border: 1px solid #000; padding: 3px 10px; margin: 14px 0 8px; text-align: center; }
.subsection-header { border: 1px solid #000; padding: 3px 10px; margin: 10px 0 6px; }
.left { text-align: left; }
.center { text-align: center; }
.box { border: 1px solid #000; padding: 0 6px; }
table { width: 100%; border-collapse: collapse; margin: 8px 0; font-size: 10pt; }
th, td { border: 1px solid #000; padding: 4px 6px; text-align: left; vertical-align: top; font-weight: normal; }
.two-column-table td { width: 50%; font-size: 9.5pt; }
p { margin: 8px 0; text-align: justify; }
.signature { text-align: center; margin-top: 90px; font-size: 9pt; line-height: 2; border-top: 1px solid #000;
             width: 60%; margin-left: 20%; padding-top: 6px; }
.wm { position: absolute; top: 45%; left: 50%; transform: translate(-50%, -50%) rotate(-50deg);
      font-size: 64pt; font-weight: bold; color: rgba(0, 0, 0, 0.07); white-space: nowrap; z-index: -1; }
"""


def find_chrome() -> str:
    for name in ("chromium", "chromium-browser", "google-chrome", "google-chrome-stable"):
        path = shutil.which(name)
        if path:
            return path
    raise SystemExit("chromium/google-chrome not found")


def build_html(md_text: str) -> str:
    pages = [p.strip() for p in md_text.split("<!--page-->")]
    body = []
    for chunk in pages:
        html = markdown.markdown(chunk, extensions=["md_in_html", "nl2br", "sane_lists"])
        body.append(f'<section class="page"><div class="wm">{WATERMARK}</div>'
                    f'<div class="hdr">{HEADER_SVG}</div>{html}</section>')
    return (f'<!DOCTYPE html><html lang="es"><head><meta charset="utf-8"><title>Oficio sintético</title>'
            f'<style>{CSS}</style></head><body>{"".join(body)}</body></html>')


def render_pdf(html: str, pdf_path: Path, chrome: str) -> None:
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "doc.html"
        src.write_text(html, encoding="utf-8")
        subprocess.run([chrome, "--headless", "--disable-gpu", "--no-sandbox", "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf_path}", src.as_uri()],
                       check=True, capture_output=True, timeout=120)


def rasterize(pdf_path: Path, dpi: int) -> list[Image.Image]:
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(["pdftoppm", "-r", str(dpi), "-png", str(pdf_path), f"{td}/p"], check=True)
        return [Image.open(p).convert("RGB") for p in sorted(Path(td).glob("p-*.png"))]


def simulate_scan(img: Image.Image, seed: int) -> Image.Image:
    """Mild, deterministic scanner look: gray paper, tiny skew, light blur and noise."""
    rng = np.random.default_rng(seed)
    angle = float(rng.uniform(-0.35, 0.35))
    img = img.rotate(angle, resample=Image.BICUBIC, expand=False, fillcolor=(255, 255, 255))
    img = img.filter(ImageFilter.GaussianBlur(0.4))
    arr = np.asarray(img).astype(np.float32)
    arr = arr * 0.96 + 6  # slightly gray paper / softer blacks
    arr += rng.normal(0, 3.0, arr.shape[:2])[..., None]
    gray = np.clip(arr.mean(axis=2, keepdims=True), 0, 255)
    return Image.fromarray(np.repeat(gray, 3, axis=2).astype(np.uint8))


def image_only_pdf(pages: list[Image.Image], pdf_path: Path, dpi: int = 300) -> None:
    jpgs = []
    for p in pages:
        buf = io.BytesIO()
        p.save(buf, "JPEG", quality=85)
        jpgs.append(Image.open(io.BytesIO(buf.getvalue())))
    jpgs[0].save(pdf_path, "PDF", resolution=dpi, save_all=True, append_images=jpgs[1:],
                 title="Oficio sintético", author="ExxerCube synthetic generator", creator="render_legacy_oficios.py")


def rebuild_docx(src: Path, dst: Path) -> None:
    """Rebuild a cover-letter DOCX from text only (fresh document: no third-party styles/metadata)."""
    import docx

    old = docx.Document(str(src))
    new = docx.Document()
    def clean(text: str) -> str:
        for old_text, new_text in DOCX_TEXT_SUBS.items():
            text = text.replace(old_text, new_text)
        return text

    tables = iter(old.tables)
    paragraphs = iter(old.paragraphs)
    for child in old.element.body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            new.add_paragraph(clean(next(paragraphs).text))
        elif tag == "tbl":
            t = next(tables)
            nt = new.add_table(rows=len(t.rows), cols=len(t.columns))
            for r, row in enumerate(t.rows):
                for c, cell in enumerate(row.cells[: len(t.columns)]):
                    nt.cell(r, c).text = clean(cell.text)
    cp = new.core_properties
    cp.author = "ExxerCube synthetic generator"
    cp.last_modified_by = "ExxerCube synthetic generator"
    cp.title = "Oficio sintético"
    cp.comments = "Synthetic test fixture"
    new.save(str(dst))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--docx-src", type=Path, help="folder holding the previous <case>.docx files to rebuild")
    ap.add_argument("--tesseract", help="tesseract binary for the 222AAA page-1 .ocr.txt")
    ap.add_argument("--tessdata", help="TESSDATA_PREFIX for --tesseract")
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    chrome = find_chrome()

    for seed, (case, (scanned, rasters)) in enumerate(CASES.items(), start=7):
        md = (HERE / f"{case}.md").read_text(encoding="utf-8")
        vector_pdf = args.out / f"{case}.vector.pdf"
        render_pdf(build_html(md), vector_pdf, chrome)
        pages = rasterize(vector_pdf, 300)
        if scanned:
            pages = [simulate_scan(p, seed * 100 + i) for i, p in enumerate(pages)]
            image_only_pdf(pages, args.out / f"{case}.pdf")
            vector_pdf.unlink()
        else:
            vector_pdf.replace(args.out / f"{case}.pdf")

        for name, w, h, fmt in rasters:
            m = re.search(r"_page-?0*(\d+)", name)
            idx = int(m.group(1)) - 1 if m else 0
            img = pages[idx].resize((w, h), Image.LANCZOS)
            kwargs = {"quality": 90, "dpi": (w * 72 // 612,) * 2} if fmt == "JPEG" else {"dpi": (w * 72 // 612,) * 2}
            img.save(args.out / name, fmt, **kwargs)

        if args.docx_src and (args.docx_src / f"{case}.docx").exists():
            rebuild_docx(args.docx_src / f"{case}.docx", args.out / f"{case}.docx")
        print(f"rendered {case}: {len(pages)} pages")

    if args.tesseract:
        env = {"TESSDATA_PREFIX": args.tessdata} if args.tessdata else None
        base = args.out / "222AAA-44444444442025_page-0001"
        subprocess.run([args.tesseract, f"{base}.jpg", f"{base}.ocr", "-l", "spa"], check=True,
                       capture_output=True, env=env)


if __name__ == "__main__":
    main()
